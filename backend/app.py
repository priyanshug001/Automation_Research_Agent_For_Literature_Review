"""Flask API for the AI Research Agent. Academic records always originate from provider responses."""
import json
import os
import re
import time
import hashlib
from datetime import datetime, timezone
from urllib.parse import quote

import requests
from flask import Flask, jsonify, request, send_file
from flask_cors import CORS
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import UniqueConstraint
from io import BytesIO
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

db = SQLAlchemy()

class Project(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(200), nullable=False)
    topic = db.Column(db.Text, nullable=False)
    question = db.Column(db.Text, default="")
    status = db.Column(db.String(40), default="created", nullable=False)
    created_at = db.Column(db.DateTime, default=lambda: datetime.now(timezone.utc), nullable=False)
    papers = db.relationship("ProjectPaper", back_populates="project", cascade="all, delete-orphan")
    reviews = db.relationship("Review", back_populates="project", cascade="all, delete-orphan")

class Paper(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    doi = db.Column(db.String(500), unique=True, nullable=True, index=True)
    title = db.Column(db.Text, nullable=False)
    authors_json = db.Column(db.Text, default="[]")
    year = db.Column(db.Integer, nullable=True)
    abstract = db.Column(db.Text, default="")
    venue = db.Column(db.Text, default="")
    source = db.Column(db.String(80), default="")
    url = db.Column(db.Text, default="")
    open_access_url = db.Column(db.Text, default="")
    metadata_status = db.Column(db.String(40), default="UNVERIFIED")
    created_at = db.Column(db.DateTime, default=lambda: datetime.now(timezone.utc), nullable=False)

    @property
    def authors(self):
        try: return json.loads(self.authors_json or "[]")
        except (ValueError, TypeError): return []

    def as_dict(self):
        return {"id": self.id, "doi": self.doi, "title": self.title, "authors": self.authors,
                "year": self.year, "abstract": self.abstract, "venue": self.venue, "source": self.source,
                "url": self.url, "open_access_url": self.open_access_url, "metadata_status": self.metadata_status}

class ProjectPaper(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    project_id = db.Column(db.Integer, db.ForeignKey("project.id", ondelete="CASCADE"), nullable=False)
    paper_id = db.Column(db.Integer, db.ForeignKey("paper.id", ondelete="CASCADE"), nullable=False)
    relevance_score = db.Column(db.Float, default=0)
    decision = db.Column(db.String(30), default="NEEDS_REVIEW")
    reason = db.Column(db.Text, default="")
    project = db.relationship("Project", back_populates="papers")
    paper = db.relationship("Paper")
    __table_args__ = (UniqueConstraint("project_id", "paper_id", name="uq_project_paper"),)

class Review(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    project_id = db.Column(db.Integer, db.ForeignKey("project.id", ondelete="CASCADE"), nullable=False)
    content = db.Column(db.Text, nullable=False)
    citation_style = db.Column(db.String(20), default="APA 7")
    created_at = db.Column(db.DateTime, default=lambda: datetime.now(timezone.utc), nullable=False)
    project = db.relationship("Project", back_populates="reviews")

def normalize_doi(value):
    if not value: return None
    value = re.sub(r"^https?://(dx\.)?doi\.org/", "", str(value).strip(), flags=re.I)
    value = re.sub(r"^doi:\s*", "", value, flags=re.I).strip().lower()
    return value or None

def normalize_title(value):
    return re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).strip()

def dedupe_key(record):
    doi = normalize_doi(record.get("doi"))
    return "doi:" + doi if doi else "title:" + normalize_title(record.get("title"))

def score_relevance(topic, paper):
    terms = set(normalize_title(topic).split())
    title_terms = set(normalize_title(paper.get("title")).split())
    abstract_terms = set(normalize_title(paper.get("abstract")).split())
    if not terms: return 0.0
    title = len(terms & title_terms) / len(terms)
    abstract = len(terms & abstract_terms) / len(terms)
    return round(min(1.0, title * 0.7 + abstract * 0.3), 3)

def _crossref(query, limit):
    mailto = os.getenv("CROSSREF_MAILTO", "")
    params = {"query.bibliographic": query, "rows": limit}
    if mailto: params["mailto"] = mailto
    r = requests.get("https://api.crossref.org/works", params=params, timeout=18,
                     headers={"User-Agent": "AIResearchAgent/1.0" + (f" (mailto:{mailto})" if mailto else "")})
    r.raise_for_status()
    items = r.json().get("message", {}).get("items", [])
    out = []
    for x in items:
        title = (x.get("title") or [""])[0]
        if not title: continue
        date_parts = x.get("published", {}).get("date-parts", [[]])[0]
        authors = [(" ".join(filter(None, [a.get("given"), a.get("family")]))).strip()
                   for a in x.get("author", [])]
        out.append({"title": title, "doi": x.get("DOI"), "authors": [a for a in authors if a],
                    "year": date_parts[0] if date_parts else None, "abstract": re.sub("<[^>]+>", " ", x.get("abstract", "")),
                    "venue": (x.get("container-title") or [""])[0], "source": "Crossref",
                    "url": x.get("URL", ""), "open_access_url": ""})
    return out

def _openalex(query, limit):
    key = os.getenv("OPENALEX_API_KEY", "")
    params = {"search": query, "per-page": min(limit, 100)}
    if key: params["api_key"] = key
    r = requests.get("https://api.openalex.org/works", params=params, timeout=18)
    r.raise_for_status()
    out = []
    for x in r.json().get("results", []):
        title = x.get("display_name") or ""
        if not title: continue
        doi = normalize_doi(x.get("doi"))
        authors = [a.get("author", {}).get("display_name", "") for a in x.get("authorships", [])]
        loc = x.get("best_oa_location") or {}
        out.append({"title": title, "doi": doi, "authors": [a for a in authors if a],
                    "year": x.get("publication_year"), "abstract": "", "venue": (x.get("primary_location") or {}).get("source", {}).get("display_name", ""),
                    "source": "OpenAlex", "url": x.get("doi") or x.get("id", ""), "open_access_url": loc.get("pdf_url") or ""})
    return out

def _arxiv(query, limit):
    if os.getenv("ARXIV_ENABLED", "true").lower() not in ("1", "true", "yes"): return []
    import xml.etree.ElementTree as ET
    r = requests.get("https://export.arxiv.org/api/query", params={"search_query": "all:" + query, "start": 0, "max_results": min(limit, 50)}, timeout=20)
    r.raise_for_status()
    root = ET.fromstring(r.text)
    ns = {"a": "http://www.w3.org/2005/Atom"}
    out = []
    for e in root.findall("a:entry", ns):
        title = " ".join((e.findtext("a:title", default="", namespaces=ns)).split())
        if not title: continue
        authors = [a.findtext("a:name", default="", namespaces=ns) for a in e.findall("a:author", ns)]
        link = next((l.attrib.get("href", "") for l in e.findall("a:link", ns) if l.attrib.get("type") == "application/pdf"), "")
        published = e.findtext("a:published", default="", namespaces=ns)
        out.append({"title": title, "doi": None, "authors": authors, "year": int(published[:4]) if published[:4].isdigit() else None,
                    "abstract": e.findtext("a:summary", default="", namespaces=ns).strip(), "venue": "arXiv",
                    "source": "arXiv", "url": e.findtext("a:id", default="", namespaces=ns), "open_access_url": link})
    return out

def _persist_paper(raw):
    doi = normalize_doi(raw.get("doi"))
    paper = Paper.query.filter_by(doi=doi).first() if doi else None
    if paper is None:
        title_key = normalize_title(raw.get("title"))
        # Title matching only for exact normalized matches; similar titles are not merged.
        if title_key:
            for candidate in Paper.query.filter(Paper.title.is_not(None)).limit(3000).all():
                if normalize_title(candidate.title) == title_key:
                    paper = candidate
                    break
    if paper is None:
        paper = Paper(doi=doi, title=raw.get("title") or "Untitled record", authors_json=json.dumps(raw.get("authors") or []),
                      year=raw.get("year"), abstract=raw.get("abstract") or "", venue=raw.get("venue") or "",
                      source=raw.get("source") or "", url=raw.get("url") or "", open_access_url=raw.get("open_access_url") or "",
                      metadata_status="METADATA_PARTIALLY_VERIFIED" if doi else "METADATA_UNVERIFIED")
        db.session.add(paper)
        db.session.flush()
    else:
        if doi and not paper.doi: paper.doi = doi
        if not paper.abstract and raw.get("abstract"): paper.abstract = raw["abstract"]
        if not paper.url and raw.get("url"): paper.url = raw["url"]
    return paper

def _llm(prompt):
    api_key = os.getenv("OPENAI_API_KEY", "")
    if not api_key: raise RuntimeError("LLM is not configured. Set OPENAI_API_KEY in backend/.env to enable generation.")
    response = requests.post("https://api.openai.com/v1/chat/completions",
        headers={"Authorization": "Bearer " + api_key, "Content-Type": "application/json"},
        json={"model": os.getenv("LLM_MODEL", "gpt-4o-mini"), "temperature": 0.2,
              "messages": [{"role": "system", "content": "You are an academic research assistant. Use only the supplied records. Never invent facts, citations, methods, results, DOIs, or novelty. State evidence limitations explicitly."},
                           {"role": "user", "content": prompt}]}, timeout=90)
    response.raise_for_status()
    return response.json()["choices"][0]["message"]["content"]

def create_app(test_config=None):
    app = Flask(__name__, instance_relative_config=True)
    app.config.update(SECRET_KEY=os.getenv("SECRET_KEY", "local-only-change-me"),
                      SQLALCHEMY_DATABASE_URI=os.getenv("DATABASE_URL", "sqlite:///research_agent.db"),
                      SQLALCHEMY_TRACK_MODIFICATIONS=False,
                      MAX_CONTENT_LENGTH=25 * 1024 * 1024)
    if test_config: app.config.update(test_config)
    CORS(app, resources={r"/api/*": {"origins": os.getenv("FRONTEND_ORIGIN", "http://localhost:5173")}})
    db.init_app(app)
    with app.app_context(): db.create_all()

    @app.errorhandler(413)
    def too_large(_): return jsonify(error="Upload exceeds the 25 MB limit."), 413

    @app.get("/api/v1/health")
    def health(): return jsonify(status="ok", database="connected", llm_configured=bool(os.getenv("OPENAI_API_KEY")))

    @app.get("/api/v1/projects")
    def list_projects():
        return jsonify(projects=[{"id": p.id, "name": p.name, "topic": p.topic, "question": p.question,
                                  "status": p.status, "paper_count": len(p.papers), "created_at": p.created_at.isoformat()} for p in Project.query.order_by(Project.created_at.desc()).all()])

    @app.post("/api/v1/projects")
    def create_project():
        data = request.get_json(silent=True) or {}
        topic = str(data.get("topic", "")).strip()
        if not topic: return jsonify(error="topic is required"), 400
        if len(topic) > 2000: return jsonify(error="topic must be 2000 characters or fewer"), 400
        p = Project(name=str(data.get("name") or topic[:120]), topic=topic, question=str(data.get("question") or ""))
        db.session.add(p); db.session.commit()
        return jsonify(project={"id": p.id, "name": p.name, "topic": p.topic, "question": p.question, "status": p.status}), 201

    @app.get("/api/v1/projects/<int:pid>")
    def get_project(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        return jsonify(project={"id": p.id, "name": p.name, "topic": p.topic, "question": p.question, "status": p.status,
                                "papers": [{"decision": pp.decision, "relevance_score": pp.relevance_score, "paper": pp.paper.as_dict()} for pp in p.papers]})

    @app.patch("/api/v1/projects/<int:pid>")
    def update_project(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        data = request.get_json(silent=True) or {}
        if "name" in data: p.name = str(data["name"])[:200]
        if "topic" in data and str(data["topic"]).strip(): p.topic = str(data["topic"]).strip()
        if "question" in data: p.question = str(data["question"])
        db.session.commit()
        return jsonify(project={"id": p.id, "name": p.name, "topic": p.topic, "question": p.question, "status": p.status})

    @app.delete("/api/v1/projects/<int:pid>")
    def delete_project(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        db.session.delete(p); db.session.commit()
        return jsonify(deleted=True)

    @app.post("/api/v1/projects/<int:pid>/search")
    def search_project(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        data = request.get_json(silent=True) or {}
        query = str(data.get("query") or p.topic).strip()
        if not query: return jsonify(error="Search query is required"), 400
        try: limit = max(1, min(int(data.get("limit", 20)), int(os.getenv("MAX_PAPERS", "50")), 50))
        except (ValueError, TypeError): return jsonify(error="limit must be an integer"), 400
        sources = data.get("sources") or ["Crossref", "OpenAlex", "arXiv"]
        adapters = {"Crossref": _crossref, "OpenAlex": _openalex, "arXiv": _arxiv}
        source_status, records = {}, []
        for source in sources:
            adapter = adapters.get(source)
            if not adapter:
                source_status[source] = {"status": "unsupported", "error": "Unknown source"}
                continue
            try:
                result = adapter(query, limit)
                records.extend(result)
                source_status[source] = {"status": "ok", "result_count": len(result)}
            except Exception as exc:
                source_status[source] = {"status": "error", "error": f"{type(exc).__name__}: provider request failed"}
        unique = {}
        for raw in records:
            key = dedupe_key(raw)
            if key not in unique or (not unique[key].get("abstract") and raw.get("abstract")): unique[key] = raw
        added, attached = 0, 0
        for raw in unique.values():
            paper = _persist_paper(raw)
            existing = ProjectPaper.query.filter_by(project_id=p.id, paper_id=paper.id).first()
            if not existing:
                db.session.add(ProjectPaper(project_id=p.id, paper_id=paper.id, relevance_score=score_relevance(p.topic, paper.as_dict()),
                                            decision="NEEDS_REVIEW", reason="Ranked by lexical overlap; this score is not proof of evidence support."))
                attached += 1
            added += 1
        p.status = "papers_discovered" if added else p.status
        db.session.commit()
        return jsonify(project_id=p.id, query=query, discovered=len(records), unique_records=len(unique), attached_to_project=attached,
                       sources=source_status, papers=[pp.paper.as_dict() | {"relevance_score": pp.relevance_score, "decision": pp.decision}
                       for pp in ProjectPaper.query.filter_by(project_id=p.id).all()]), 200

    @app.get("/api/v1/projects/<int:pid>/papers")
    def list_papers(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        return jsonify(papers=[pp.paper.as_dict() | {"relevance_score": pp.relevance_score, "decision": pp.decision, "reason": pp.reason}
                              for pp in p.papers])

    @app.patch("/api/v1/projects/<int:pid>/papers/<int:paper_id>/decision")
    def decision(pid, paper_id):
        pp = ProjectPaper.query.filter_by(project_id=pid, paper_id=paper_id).first()
        if not pp: return jsonify(error="Paper not found in this project"), 404
        data = request.get_json(silent=True) or {}
        value = data.get("decision")
        if value not in ("INCLUDED", "EXCLUDED", "NEEDS_REVIEW"): return jsonify(error="decision must be INCLUDED, EXCLUDED, or NEEDS_REVIEW"), 400
        pp.decision = value; pp.reason = str(data.get("reason") or "User decision")
        db.session.commit()
        return jsonify(decision=value, paper_id=paper_id)

    @app.post("/api/v1/projects/<int:pid>/review/generate")
    def generate_review(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        included = [pp.paper.as_dict() for pp in p.papers if pp.decision == "INCLUDED"]
        if not included: return jsonify(error="Include at least one paper before generating a review."), 400
        payload = {"topic": p.topic, "research_question": p.question, "papers": included}
        try:
            content = _llm("Write a structured literature review grounded strictly in these records. Cite studies inline by their exact titles and include a References section with the provided DOI/source URLs. Do not imply full-text reading when only abstracts/metadata are provided. Flag missing evidence and limits.\n" + json.dumps(payload, ensure_ascii=False))
        except RuntimeError as exc: return jsonify(error=str(exc)), 503
        except requests.RequestException: return jsonify(error="LLM provider request failed. Check configuration, quota, and connectivity."), 502
        review = Review(project_id=p.id, content=content, citation_style=(request.get_json(silent=True) or {}).get("citation_style", "APA 7"))
        db.session.add(review); p.status = "review_generated"; db.session.commit()
        return jsonify(review={"id": review.id, "content": review.content, "citation_style": review.citation_style}), 201

    @app.get("/api/v1/projects/<int:pid>/reviews")
    def list_reviews(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        return jsonify(reviews=[{"id": r.id, "content": r.content, "citation_style": r.citation_style, "created_at": r.created_at.isoformat()} for r in p.reviews])

    @app.post("/api/v1/projects/<int:pid>/audit")
    def audit(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        reviews = [{"id": r.id, "cited_titles_found": [pp.paper.title for pp in p.papers if pp.paper.title.lower() in r.content.lower()],
                    "included_papers": sum(pp.decision == "INCLUDED" for pp in p.papers),
                    "note": "Title-string matching is a heuristic audit, not semantic claim-support verification."} for r in p.reviews]
        return jsonify(reviews=reviews)

    @app.post("/api/v1/projects/<int:pid>/export")
    def export_review(pid):
        p = db.session.get(Project, pid)
        if not p: return jsonify(error="Project not found"), 404
        review = Review.query.filter_by(project_id=pid).order_by(Review.created_at.desc()).first()
        if not review: return jsonify(error="Generate a review before exporting."), 400
        fmt = request.args.get("format", "docx").lower()
        buf = BytesIO()
        if fmt == "docx":
            doc = Document(); doc.add_heading(p.name, 0); doc.add_paragraph("Research topic: " + p.topic)
            doc.add_paragraph("Research question: " + (p.question or "Not specified"))
            for block in review.content.split("\n"):
                if block.strip(): doc.add_paragraph(block)
            doc.save(buf); mimetype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"; filename = "literature-review.docx"
        elif fmt == "pdf":
            c = canvas.Canvas(buf, pagesize=letter); width, height = letter; y = height - 54
            for line in (p.name + "\nTopic: " + p.topic + "\n\n" + review.content).splitlines():
                from reportlab.pdfbase.pdfmetrics import stringWidth
                words, current = line.split(), ""
                for word in words:
                    if stringWidth(current + " " + word, "Helvetica", 10) > width - 100:
                        c.drawString(50, y, current); y -= 14; current = word
                    else: current = (current + " " + word).strip()
                if current: c.drawString(50, y, current); y -= 14
                if y < 60: c.showPage(); y = height - 54
            c.save(); mimetype = "application/pdf"; filename = "literature-review.pdf"
        else: return jsonify(error="format must be docx or pdf"), 400
        buf.seek(0)
        return send_file(buf, mimetype=mimetype, as_attachment=True, download_name=filename)

    return app
